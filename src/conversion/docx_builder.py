"""Сборка DOCX из распознанной разметки страниц."""

from __future__ import annotations

import io
import statistics
from dataclasses import dataclass, field

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt

from .layout import Block, Paragraph, Row, Span, Table, parse_inline
from .pages import PageGeometry, estimate_font_size

FONT_NAME = "Times New Roman"
DEFAULT_FONT_SIZE = 12.0
MIN_FONT_SIZE, MAX_FONT_SIZE = 10.0, 14.0
FIRST_LINE_INDENT = Cm(1.25)
# Поля по умолчанию и допустимые пределы измеренных полей, см.
DEFAULT_MARGINS_CM = (3.0, 1.5, 2.0, 2.0)
MARGIN_LIMITS_CM = ((1.0, 4.0), (0.8, 3.0), (1.0, 3.0))
BOTTOM_MARGIN_CM = 1.5

_ALIGNMENT = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
_SIZE_STEP = {"small": -2.0, "normal": 0.0, "large": 2.0, "xlarge": 4.0}


@dataclass
class PageContent:
    number: int
    geometry: PageGeometry
    blocks: list[Block] = field(default_factory=list)
    failed: bool = False


def build_docx(pages: list[PageContent]) -> bytes:
    document = Document()
    base_size = _body_font_size(pages)
    _setup_defaults(document, base_size, _line_spacing(pages, base_size))
    builder = _Builder(document, base_size)
    for index, page in enumerate(pages):
        section = document.sections[0] if index == 0 else _new_page_section(document)
        _setup_section(section, page.geometry)
        builder.text_width = section.page_width - section.left_margin - section.right_margin
        builder.add_page(page)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class _Builder:
    def __init__(self, document, base_size: float) -> None:
        self.document = document
        self.base_size = base_size
        self.text_width = Emu(0)
        self.last_was_table = False

    def add_page(self, page: PageContent) -> None:
        self.last_was_table = False
        if page.failed:
            paragraph = self.document.add_paragraph()
            run = paragraph.add_run(f"[Страница {page.number} не распознана]")
            run.italic = True
            return
        for block in page.blocks:
            if isinstance(block, Paragraph):
                self._paragraph(block)
            elif isinstance(block, Row):
                self._row(block)
            else:
                self._table(block)

    def _paragraph(self, block: Paragraph) -> None:
        paragraph = self._new_paragraph(block.gap)
        paragraph.alignment = _ALIGNMENT[block.align]
        if block.first_line:
            paragraph.paragraph_format.first_line_indent = FIRST_LINE_INDENT
        if block.indent:
            paragraph.paragraph_format.left_indent = Emu(int(self.text_width * block.indent))
        self._add_text(paragraph, block.text, block.size, block.bold)

    def _row(self, block: Row) -> None:
        if any("\n" in cell for cell in block.cells):
            # Колонки в несколько строк — таблица без рамок.
            table = Table([block.cells], borders=False, size=block.size, gap=block.gap)
            self._table(table, bold=block.bold, row_alignment=True)
            return

        paragraph = self._new_paragraph(block.gap)
        tab_stops = paragraph.paragraph_format.tab_stops
        if len(block.cells) == 3:
            tab_stops.add_tab_stop(Emu(int(self.text_width / 2)), WD_TAB_ALIGNMENT.CENTER)
        tab_stops.add_tab_stop(self.text_width, WD_TAB_ALIGNMENT.RIGHT)
        for index, cell in enumerate(block.cells):
            if index:
                paragraph.add_run("\t")
            self._add_text(paragraph, cell, block.size, block.bold)

    def _table(self, block: Table, bold: bool = False, row_alignment: bool = False) -> None:
        if self.last_was_table:
            # Соседние таблицы Word склеивает в одну.
            self.document.add_paragraph()
        elif block.gap:
            self.document.add_paragraph()

        columns = len(block.rows[0])
        table = self.document.add_table(rows=len(block.rows), cols=columns)
        if block.borders:
            table.style = self.document.styles["Table Grid"]
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        widths = [Emu(int(self.text_width * share)) for share in block.widths or _column_widths(block.rows)]
        # Word берёт ширину из ячеек, LibreOffice — из сетки таблицы.
        for column, width in zip(table.columns, widths):
            column.width = width
        for row_cells, row in zip((r.cells for r in table.rows), block.rows):
            for column, (cell, text) in enumerate(zip(row_cells, row)):
                cell.width = widths[column]
                paragraph = cell.paragraphs[0]
                if row_alignment:
                    paragraph.alignment = _row_cell_alignment(column, columns)
                self._add_text(paragraph, text, block.size, bold)
        self.last_was_table = True

    def _new_paragraph(self, gap: bool):
        self.last_was_table = False
        paragraph = self.document.add_paragraph()
        if gap:
            paragraph.paragraph_format.space_before = Pt(round(self.base_size))
        return paragraph

    def _add_text(self, paragraph, text: str, size: str, bold: bool) -> None:
        font_size = Pt(max(self.base_size + _SIZE_STEP[size], 7)) if size != "normal" else None
        for span in parse_inline(text):
            for index, line in enumerate(span.text.split("\n")):
                run = paragraph.add_run()
                if index:
                    run.add_break()
                run.add_text(line)
                _style_run(run, span, bold, font_size)


def _style_run(run, span: Span, bold: bool, font_size) -> None:
    if bold or span.bold:
        run.bold = True
    if span.italic:
        run.italic = True
    if span.underline:
        run.underline = True
    if font_size is not None:
        run.font.size = font_size


def _row_cell_alignment(column: int, columns: int):
    if column == 0:
        return WD_ALIGN_PARAGRAPH.LEFT
    if column == columns - 1:
        return WD_ALIGN_PARAGRAPH.RIGHT
    return WD_ALIGN_PARAGRAPH.CENTER


def _column_widths(rows: list[list[str]]) -> list[float]:
    # Ширина по самому длинному тексту в колонке, но не уже 8 % строки.
    longest = [max(len(row[c]) for row in rows) for c in range(len(rows[0]))]
    weights = [max(length, 1) ** 0.75 for length in longest]
    total = sum(weights)
    shares = [max(weight / total, 0.08) for weight in weights]
    total = sum(shares)
    return [share / total for share in shares]


def _body_font_size(pages: list[PageContent]) -> float:
    x_heights = [p.geometry.x_height_pt for p in pages if p.geometry.x_height_pt]
    if not x_heights:
        return DEFAULT_FONT_SIZE
    size = estimate_font_size(statistics.median(x_heights))
    return min(max(round(size * 2) / 2, MIN_FONT_SIZE), MAX_FONT_SIZE)


def _line_spacing(pages: list[PageContent], base_size: float) -> float:
    # Одинарный интервал Times New Roman — около 1,15 кегля; разреженнее 1,15
    # не делаем, чтобы текст страницы не переползал на следующую.
    pitches = [p.geometry.line_pitch_pt for p in pages if p.geometry.line_pitch_pt]
    if not pitches:
        return 1.0
    ratio = statistics.median(pitches) / (1.15 * base_size)
    return round(min(max(ratio, 1.0), 1.15), 2)


def _setup_defaults(document, base_size: float, line_spacing: float) -> None:
    """Шрифт и интервалы по умолчанию вместо Calibri и 10 пт после абзаца из шаблона."""
    styles = document.styles.element
    defaults = styles.find(qn("w:docDefaults"))
    run_props = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr"))
    fonts = run_props.find(qn("w:rFonts"))
    for attribute in list(fonts.attrib):
        del fonts.attrib[attribute]
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attribute), FONT_NAME)
    half_points = str(int(round(base_size * 2)))
    for tag in ("w:sz", "w:szCs"):
        run_props.find(qn(tag)).set(qn("w:val"), half_points)

    spacing = defaults.find(qn("w:pPrDefault")).find(qn("w:pPr")).find(qn("w:spacing"))
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), str(int(round(240 * line_spacing))))
    spacing.set(qn("w:lineRule"), "auto")


def _setup_section(section, geometry: PageGeometry) -> None:
    landscape = geometry.width_pt > geometry.height_pt
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Pt(geometry.width_pt)
    section.page_height = Pt(geometry.height_pt)

    left, right, top, bottom = DEFAULT_MARGINS_CM
    if geometry.margins_pt:
        measured = [value / 28.35 for value in geometry.margins_pt[:3]]
        left, right, top = (
            min(max(value, low), high) for value, (low, high) in zip(measured, MARGIN_LIMITS_CM)
        )
        # Снизу текст часто кончается раньше конца листа, мерить там нечего.
        bottom = BOTTOM_MARGIN_CM
    section.left_margin = Cm(left)
    section.right_margin = Cm(right)
    section.top_margin = Cm(top)
    section.bottom_margin = Cm(bottom)
    section.header_distance = Cm(min(top, 1.25) / 2)
    section.footer_distance = Cm(bottom / 2)


def _new_page_section(document):
    """Новый раздел с новой страницы без пустого абзаца под разрыв.

    python-docx кладёт разрыв раздела в отдельный пустой абзац; на полностью
    занятой странице он уезжает на следующую и даёт пустой лист. Поэтому
    разрыв переносится в последний абзац страницы, если он есть.
    """
    section = document.add_section(WD_SECTION.NEW_PAGE)
    body = document.element.body
    break_paragraph = body.findall(qn("w:p"))[-1]
    previous = break_paragraph.getprevious()
    if previous is not None and previous.tag == qn("w:p"):
        properties = previous.get_or_add_pPr()
        if properties.sectPr is None:
            section_properties = break_paragraph.pPr.sectPr
            properties._insert_sectPr(section_properties)
            body.remove(break_paragraph)
    return section
