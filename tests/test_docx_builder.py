import io

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from src.conversion.docx_builder import PageContent, build_docx
from src.conversion.layout import Paragraph, Row, Table
from src.conversion.pages import PageGeometry

A4 = (595.28, 841.89)
A5 = (419.53, 595.28)


def _build(*pages: PageContent) -> Document:
    return Document(io.BytesIO(build_docx(list(pages))))


def _page(number, blocks, size=A4, **geometry):
    return PageContent(number, PageGeometry(*size, **geometry), blocks)


def test_each_page_is_a_section_with_its_own_size():
    doc = _build(
        _page(1, [Paragraph("small page")], size=A5),
        _page(2, [Paragraph("portrait")]),
        _page(3, [Paragraph("landscape")], size=(A4[1], A4[0])),
    )
    sizes = [(round(s.page_width.pt), round(s.page_height.pt), s.orientation) for s in doc.sections]
    assert sizes == [
        (420, 595, WD_ORIENT.PORTRAIT),
        (595, 842, WD_ORIENT.PORTRAIT),
        (842, 595, WD_ORIENT.LANDSCAPE),
    ]
    # Разрывы разделов живут в последних абзацах страниц, пустых абзацев нет.
    assert [p.text for p in doc.paragraphs] == ["small page", "portrait", "landscape"]


def test_measured_margins_are_clamped():
    doc = _build(_page(1, [Paragraph("x")], margins_pt=(3.17 * 28.35, 0.1 * 28.35, 9 * 28.35, 0)))
    section = doc.sections[0]
    assert round(section.left_margin.cm, 2) == 3.17
    assert round(section.right_margin.cm, 2) == 0.8
    assert round(section.top_margin.cm, 2) == 3.0
    assert round(section.bottom_margin.cm, 2) == 1.5


def test_defaults_use_times_new_roman_and_no_spacing_after():
    doc = _build(_page(1, [Paragraph("x")], x_height_pt=6.3, line_pitch_pt=20.0))
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    fonts = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr")).find(qn("w:rFonts"))
    assert fonts.get(qn("w:ascii")) == "Times New Roman"
    assert fonts.get(qn("w:asciiTheme")) is None
    size = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr")).find(qn("w:sz"))
    assert size.get(qn("w:val")) == "28"  # 14 пт по высоте строчных 6,3 пт
    spacing = defaults.find(qn("w:pPrDefault")).find(qn("w:pPr")).find(qn("w:spacing"))
    assert spacing.get(qn("w:after")) == "0"


def test_paragraph_formatting_and_inline_markup():
    doc = _build(
        _page(
            1,
            [
                Paragraph("**4.** Raisi (**S. Rasulev**):", align="justify", first_line=True),
                Paragraph("Title\nsecond line", align="center", size="large", bold=True, gap=True),
                Paragraph("ilova", indent=0.5),
            ],
        )
    )
    first, title, appendix = doc.paragraphs
    assert first.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert abs(first.paragraph_format.first_line_indent - Cm(1.25)) < Cm(0.01)
    assert [(r.text, bool(r.bold)) for r in first.runs] == [
        ("4.", True),
        (" Raisi (", False),
        ("S. Rasulev", True),
        ("):", False),
    ]

    assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert title.paragraph_format.space_before is not None
    assert all(r.bold and r.font.size == Pt(14) for r in title.runs)
    breaks = title._p.findall(".//" + qn("w:br"))
    assert len(breaks) == 1 and breaks[0].get(qn("w:type")) in (None, str(WD_BREAK.LINE))
    assert title.text.replace("\n", " ") == "Title second line"

    section = doc.sections[0]
    text_width = section.page_width - section.left_margin - section.right_margin
    assert abs(appendix.paragraph_format.left_indent - text_width / 2) < 1000


def test_row_uses_right_tab_stop():
    doc = _build(_page(1, [Row(["Direktor", "M.Olloyorov"], bold=True)]))
    paragraph = doc.paragraphs[0]
    assert paragraph.text == "Direktor\tM.Olloyorov"
    stops = list(paragraph.paragraph_format.tab_stops)
    section = doc.sections[0]
    assert stops[-1].position == section.page_width - section.left_margin - section.right_margin


def test_multiline_row_becomes_borderless_table():
    doc = _build(_page(1, [Row(["№83/3B\n2026-yil", "A.S.MUKIMOVA\nDirektor"])]))
    table = doc.tables[0]
    assert table.style.name != "Table Grid"
    left, right = table.rows[0].cells
    assert left.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert right.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT


def test_tables_get_borders_widths_and_separator():
    doc = _build(
        _page(
            1,
            [
                Table([["1.", "*Komissiya raisi*"]], borders=False, widths=[0.2, 0.8]),
                Table([["a", "b"]]),
            ],
        )
    )
    first, second = doc.tables
    assert first.style.name != "Table Grid"
    assert second.style.name == "Table Grid"
    section = doc.sections[0]
    text_width = section.page_width - section.left_margin - section.right_margin
    assert abs(first.columns[1].width - text_width * 0.8) < 1000
    assert first.rows[0].cells[1].paragraphs[0].runs[0].italic
    # Между таблицами есть абзац, иначе Word склеит их в одну.
    body = [child.tag.split("}")[1] for child in doc.element.body]
    assert body[:3] == ["tbl", "p", "tbl"]


def test_failed_page_gets_placeholder():
    doc = _build(_page(1, [Paragraph("ok")]), PageContent(2, PageGeometry(*A4), [], failed=True))
    assert doc.paragraphs[-1].text == "[Страница 2 не распознана]"
